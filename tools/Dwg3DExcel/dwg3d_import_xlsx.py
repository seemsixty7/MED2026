"""Import Category / Description / Note edits from the bulk-edit workbook into Dwg3D\\Dwg3DCatalog.db.

    py -3.12 dwg3d_import_xlsx.py [XLSX] [--db PATH] [--dry-run]

Rules (same as the MED3DLIB palette "Import from Excel"):
* rows are matched by File (relative path, case-insensitive; falls back to a unique file name);
* only non-blank cells are applied, and only where they differ from the DB;
* a cell containing (clear) empties that field;
* the DB is backed up to Dwg3DCatalog.db.bak-YYYYMMDD-HHMMSS before anything is written.
"""
import argparse, datetime, os, shutil, sqlite3, sys
from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
DEFAULT_DB = os.path.normpath(os.path.join(HERE, "..", "..", "Dwg3D", "Dwg3DCatalog.db"))
FIELDS = [("Category", "category"), ("Description", "description"), ("Note", "note")]


def text(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx", nargs="?", default=None)
    ap.add_argument("--db", default=DEFAULT_DB)
    ap.add_argument("--dry-run", action="store_true")
    a = ap.parse_args()
    xlsx = a.xlsx or os.path.join(os.path.dirname(a.db), "Dwg3DCatalog_BulkEdit.xlsx")
    wb = load_workbook(xlsx, read_only=True, data_only=True)
    ws = wb["Blocks"] if "Blocks" in wb.sheetnames else wb.worksheets[0]
    it = ws.iter_rows(values_only=True)
    header = [text(h).strip().lower() for h in next(it)]
    if "file" not in header:
        sys.exit("No 'File' column in the first row of sheet " + ws.title)
    fcol = header.index("file")
    cols = [(name, dbcol, header.index(name.lower())) for name, dbcol in FIELDS if name.lower() in header]

    con = sqlite3.connect(a.db)
    db = {r[1].lower(): r for r in con.execute("SELECT id,file,category,description,note FROM blocks")}
    by_name = {}
    for r in db.values():
        by_name.setdefault(os.path.basename(r[1]).lower(), []).append(r)
    idx = {"category": 2, "description": 3, "note": 4}

    rows = matched = 0
    not_found = []
    changes = []  # (id, file, dbcol, old, new)
    for row in it:
        f = text(row[fcol] if fcol < len(row) else None).strip()
        if not f:
            continue
        rows += 1
        r = db.get(f.lower())
        if r is None:
            cand = by_name.get(os.path.basename(f).lower(), [])
            r = cand[0] if len(cand) == 1 else None
        if r is None:
            not_found.append(f)
            continue
        matched += 1
        for name, dbcol, c in cols:
            v = text(row[c] if c < len(row) else None)
            v = v.rstrip() if dbcol == "note" else v.strip()
            if not v:
                continue
            if v.lower() == "(clear)":
                v = ""
            old = r[idx[dbcol]] or ""
            if v != old:
                changes.append((r[0], r[1], dbcol, old, v))
    wb.close()

    entries = len({c[0] for c in changes})
    per = {d: sum(1 for c in changes if c[2] == d) for _, d in FIELDS}
    print(f"{rows} rows read, {matched} matched, {len(not_found)} not found, {entries} blocks to update, {len(changes)} fields "
          f"(category {per['category']}, description {per['description']}, note {per['note']})")
    for f in not_found[:20]:
        print("  not found:", f)
    if a.dry_run or not changes:
        print("Dry run: nothing written." if a.dry_run else "Nothing to update.")
        return
    con.close()
    bak = a.db + ".bak-" + datetime.datetime.now().strftime("%Y%m%d-%H%M%S")
    shutil.copy2(a.db, bak)
    print("Backup:", bak)
    con = sqlite3.connect(a.db)
    now = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    with con:
        for bid, f, dbcol, old, new in changes:
            con.execute(f"UPDATE blocks SET {dbcol}=?, date_updated=? WHERE id=?", (new, now, bid))
            if dbcol == "category" and new.strip():
                con.execute("INSERT OR IGNORE INTO categories(name) VALUES(?)", (new.strip(),))
    con.close()
    print("Updated.")


if __name__ == "__main__":
    main()
